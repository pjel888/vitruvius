"""Bipartite table-to-graph extractor for Tasks B.3 / B.3.1."""

from __future__ import annotations

import csv
import hashlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_ID_LEN = 12


def _stable_id(*parts: str) -> str:
    """First 12 hex chars of SHA-1 over a NUL-joined parts tuple."""
    joined = "\x00".join(parts)
    return hashlib.sha1(joined.encode("utf-8")).hexdigest()[:_ID_LEN]


@dataclass(frozen=True)
class TableNode:
    id: str
    name: str
    kind: str


@dataclass(frozen=True)
class TableEdge:
    edge_id: str
    row_id: str
    col_id: str
    cell_value: str


@dataclass
class TableSubgraph:
    source_path: str
    row_nodes: list[TableNode] = field(default_factory=list)
    column_nodes: list[TableNode] = field(default_factory=list)
    edges: list[TableEdge] = field(default_factory=list)

    def to_json_dict(self) -> dict[str, Any]:
        return {
            "source_path": self.source_path,
            "row_nodes": [n.__dict__ for n in self.row_nodes],
            "column_nodes": [n.__dict__ for n in self.column_nodes],
            "edges": [e.__dict__ for e in self.edges],
        }


def _build_subgraph_from_rows(
    rows: list[list[str]],
    source_path: str,
    id_salt: str,
) -> TableSubgraph:
    subgraph = TableSubgraph(source_path=source_path)
    if not rows:
        return subgraph

    header = rows[0]
    data_columns = header[1:]
    for col_name in data_columns:
        subgraph.column_nodes.append(
            TableNode(
                id=f"col:{_stable_id(id_salt, col_name)}",
                name=col_name,
                kind="col",
            )
        )

    for raw_row in rows[1:]:
        if not any(cell.strip() for cell in raw_row):
            continue
        row_label = raw_row[0]
        row_id = f"row:{_stable_id(id_salt, row_label)}"
        subgraph.row_nodes.append(
            TableNode(id=row_id, name=row_label, kind="row")
        )
        for col_node, cell_value in zip(
            subgraph.column_nodes, raw_row[1 : 1 + len(data_columns)]
        ):
            subgraph.edges.append(
                TableEdge(
                    edge_id=_stable_id(
                        id_salt, row_label, col_node.name, cell_value
                    ),
                    row_id=row_id,
                    col_id=col_node.id,
                    cell_value=cell_value,
                )
            )

    return subgraph


def extract_table_subgraph(csv_path: Path) -> TableSubgraph:
    """Project a CSV into a deterministic bipartite TableSubgraph.

    Assumes the first column is the row label. Empty/whitespace-only rows
    are skipped. Empty cells become edges with cell_value == "".
    """
    with csv_path.open("r", encoding="utf-8", newline="") as handle:
        rows = [list(row) for row in csv.reader(handle)]
    return _build_subgraph_from_rows(
        rows=rows,
        source_path=str(csv_path),
        id_salt=csv_path.name,
    )


def _coerce_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def extract_xlsx_subgraphs(xlsx_path: Path) -> list[TableSubgraph]:
    """Parse every sheet of an XLSX file into a TableSubgraph.

    Returns one subgraph per sheet in sheet-order. Empty sheets yield an
    empty TableSubgraph. source_path is f"{xlsx_path.name}#{sheet_name}".
    """
    workbook = load_workbook(xlsx_path, data_only=True)
    filename = xlsx_path.name
    subgraphs: list[TableSubgraph] = []

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        source_path = f"{filename}#{sheet_name}"
        id_salt = source_path

        rows: list[list[str]] = []
        for raw_row in worksheet.iter_rows(values_only=True):
            rows.append([_coerce_cell(cell) for cell in raw_row])

        subgraphs.append(
            _build_subgraph_from_rows(
                rows=rows,
                source_path=source_path,
                id_salt=id_salt,
            )
        )

    return subgraphs

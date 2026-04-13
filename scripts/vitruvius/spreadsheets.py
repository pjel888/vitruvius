from __future__ import annotations

import csv
import json
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .constants import CONVERSION_LAYER_NAME, RAW_LAYER_NAME
from .paths import (
    build_spreadsheet_csv_relative_dir,
    build_spreadsheet_json_relative_dir,
    build_spreadsheet_manifest_relative_path,
    ensure_registry_dir,
    ensure_runs_dir,
    get_repo_root,
)
from .reports import (
    append_spreadsheet_registry_row,
    ensure_spreadsheet_registry_exists,
    snapshot_spreadsheet_registry,
    write_json,
    write_text,
)
from .utils import safe_name_for_report, utc_now_stamp


SUPPORTED_SPREADSHEET_EXTENSIONS = {
    ".xlsx",
    ".xls",
}


def is_supported_spreadsheet_file(path: str | Path) -> bool:
    return Path(path).suffix.lower() in SUPPORTED_SPREADSHEET_EXTENSIONS


def get_default_libreoffice_candidates() -> list[Path]:
    candidates = [
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        Path("/usr/bin/libreoffice"),
        Path("/snap/bin/libreoffice"),
        Path("/usr/bin/soffice"),
    ]
    return candidates


def find_soffice_executable() -> str:
    which_result = shutil.which("soffice") or shutil.which("libreoffice")
    if which_result:
        return which_result

    for candidate in get_default_libreoffice_candidates():
        if candidate.exists():
            return str(candidate)

    raise FileNotFoundError(
        "LibreOffice executable not found. Install LibreOffice and make sure "
        "soffice is available in PATH, or installed in a standard location."
    )


def sanitize_sheet_name_for_filename(sheet_name: str) -> str:
    sanitized = "".join(ch if ch.isalnum() else "_" for ch in sheet_name).strip("_")
    return sanitized or "Sheet"


def normalize_cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def count_nonempty_columns(rows: list[list[Any]]) -> int:
    max_cols = 0
    for row in rows:
        if not row:
            continue

        last_nonempty_index = 0
        for idx, value in enumerate(row, start=1):
            if value not in (None, ""):
                last_nonempty_index = idx

        if last_nonempty_index > max_cols:
            max_cols = last_nonempty_index

    return max_cols


def trim_trailing_empty_cells(row: list[Any]) -> list[Any]:
    trimmed = list(row)
    while trimmed and trimmed[-1] in (None, ""):
        trimmed.pop()
    return trimmed


def workbook_to_xlsx_if_needed(relative_source_path: str) -> dict[str, Any]:
    repo_root = get_repo_root()
    source_path = repo_root / RAW_LAYER_NAME / relative_source_path

    if not source_path.exists():
        raise FileNotFoundError(f"Spreadsheet source file not found: {source_path}")

    suffix = source_path.suffix.lower()
    if suffix not in SUPPORTED_SPREADSHEET_EXTENSIONS:
        raise ValueError(
            f"Unsupported spreadsheet file type: {suffix}. "
            f"Supported types: {sorted(SUPPORTED_SPREADSHEET_EXTENSIONS)}"
        )

    stat = source_path.stat()
    top_category = Path(relative_source_path).parts[0] if len(Path(relative_source_path).parts) > 0 else ""

    if suffix == ".xlsx":
        return {
            "source_relative_path": relative_source_path,
            "source_absolute_path": str(source_path),
            "source_extension": suffix,
            "source_size_bytes": stat.st_size,
            "source_modified_utc": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
            "top_category": top_category,
            "xlsx_workbook_path": str(source_path),
            "xlsx_workbook_relative_path": f"{RAW_LAYER_NAME}/{Path(relative_source_path).as_posix()}",
            "conversion_used": "none",
        }

    soffice = find_soffice_executable()

    xlsx_rel = build_spreadsheet_manifest_relative_path(relative_source_path).with_suffix(".xlsx")
    xlsx_abs = repo_root / CONVERSION_LAYER_NAME / xlsx_rel
    xlsx_abs.parent.mkdir(parents=True, exist_ok=True)

    if xlsx_abs.exists():
        xlsx_abs.unlink()

    command = [
        soffice,
        "--headless",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(xlsx_abs.parent),
        str(source_path),
    ]

    completed = subprocess.run(
        command,
        capture_output=True,
        text=True,
        check=False,
    )

    if completed.returncode != 0:
        raise RuntimeError(
            "LibreOffice spreadsheet conversion failed.\n"
            f"Command: {' '.join(command)}\n"
            f"STDOUT: {completed.stdout.strip()}\n"
            f"STDERR: {completed.stderr.strip()}"
        )

    expected_output = xlsx_abs.parent / f"{source_path.stem}.xlsx"

    if not expected_output.exists():
        raise FileNotFoundError(
            f"Expected converted XLSX not found after LibreOffice run: {expected_output}"
        )

    if expected_output.resolve() != xlsx_abs.resolve():
        if xlsx_abs.exists():
            xlsx_abs.unlink()
        expected_output.replace(xlsx_abs)

    return {
        "source_relative_path": relative_source_path,
        "source_absolute_path": str(source_path),
        "source_extension": suffix,
        "source_size_bytes": stat.st_size,
        "source_modified_utc": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
        "top_category": top_category,
        "xlsx_workbook_path": str(xlsx_abs),
        "xlsx_workbook_relative_path": f"{CONVERSION_LAYER_NAME}/{xlsx_rel.as_posix()}",
        "conversion_used": "xls_to_xlsx_via_libreoffice",
    }


def _write_spreadsheet_registry_entry(
    *,
    run_command: str,
    source_relative_path: str,
    logical_source_relative_path: str,
    source_extension: str,
    source_size_bytes: int | str,
    source_modified_utc: str,
    top_category: str,
    sheet_name: str,
    output_csv_relative_path: str = "",
    output_json_relative_path: str = "",
    row_count: int | str = "",
    column_count: int | str = "",
    status: str = "",
    error_message: str = "",
) -> None:
    repo_root = get_repo_root()
    registry_dir = ensure_registry_dir(repo_root)

    ensure_spreadsheet_registry_exists(registry_dir)

    row = {
        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "run_command": run_command,
        "source_relative_path": source_relative_path,
        "logical_source_relative_path": logical_source_relative_path,
        "source_extension": source_extension,
        "source_size_bytes": source_size_bytes,
        "source_modified_utc": source_modified_utc,
        "top_category": top_category,
        "sheet_name": sheet_name,
        "output_csv_relative_path": output_csv_relative_path,
        "output_json_relative_path": output_json_relative_path,
        "row_count": row_count,
        "column_count": column_count,
        "status": status,
        "error_message": error_message,
    }

    append_spreadsheet_registry_row(registry_dir, row)


def export_spreadsheet_workbook(
    relative_source_path: str,
    *,
    run_command: str = "convert-spreadsheet",
) -> dict[str, Any]:
    repo_root = get_repo_root()

    workbook_info = workbook_to_xlsx_if_needed(relative_source_path)
    workbook_path = Path(workbook_info["xlsx_workbook_path"])

    csv_dir_rel = build_spreadsheet_csv_relative_dir(relative_source_path)
    json_dir_rel = build_spreadsheet_json_relative_dir(relative_source_path)
    manifest_rel = build_spreadsheet_manifest_relative_path(relative_source_path)

    csv_dir_abs = repo_root / CONVERSION_LAYER_NAME / csv_dir_rel
    json_dir_abs = repo_root / CONVERSION_LAYER_NAME / json_dir_rel
    manifest_abs = repo_root / CONVERSION_LAYER_NAME / manifest_rel

    csv_dir_abs.mkdir(parents=True, exist_ok=True)
    json_dir_abs.mkdir(parents=True, exist_ok=True)
    manifest_abs.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(
        filename=str(workbook_path),
        data_only=True,
        read_only=True,
    )

    sheet_reports: list[dict[str, Any]] = []

    try:
        for worksheet in workbook.worksheets:
            sheet_name = worksheet.title
            safe_sheet_name = sanitize_sheet_name_for_filename(sheet_name)

            sheet_rows_raw = []
            for row in worksheet.iter_rows(values_only=True):
                normalized_row = [normalize_cell_value(value) for value in row]
                sheet_rows_raw.append(trim_trailing_empty_cells(normalized_row))

            row_count = len(sheet_rows_raw)
            column_count = count_nonempty_columns(sheet_rows_raw)

            csv_rel = csv_dir_rel / f"{safe_sheet_name}.csv"
            json_rel = json_dir_rel / f"{safe_sheet_name}.json"

            csv_abs = repo_root / CONVERSION_LAYER_NAME / csv_rel
            json_abs = repo_root / CONVERSION_LAYER_NAME / json_rel

            with csv_abs.open("w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                for row in sheet_rows_raw:
                    writer.writerow(row)

            json_payload = {
                "workbook_source_relative_path": relative_source_path,
                "sheet_name": sheet_name,
                "row_count": row_count,
                "column_count": column_count,
                "rows": sheet_rows_raw,
            }
            write_json(json_abs, json_payload)

            _write_spreadsheet_registry_entry(
                run_command=run_command,
                source_relative_path=f"{RAW_LAYER_NAME}/{Path(relative_source_path).as_posix()}",
                logical_source_relative_path=Path(relative_source_path).as_posix(),
                source_extension=workbook_info["source_extension"],
                source_size_bytes=workbook_info["source_size_bytes"],
                source_modified_utc=workbook_info["source_modified_utc"],
                top_category=workbook_info["top_category"],
                sheet_name=sheet_name,
                output_csv_relative_path=f"{CONVERSION_LAYER_NAME}/{csv_rel.as_posix()}",
                output_json_relative_path=f"{CONVERSION_LAYER_NAME}/{json_rel.as_posix()}",
                row_count=row_count,
                column_count=column_count,
                status="SUCCESS",
                error_message="",
            )

            sheet_reports.append(
                {
                    "sheet_name": sheet_name,
                    "safe_sheet_name": safe_sheet_name,
                    "row_count": row_count,
                    "column_count": column_count,
                    "output_csv_relative_path": f"{CONVERSION_LAYER_NAME}/{csv_rel.as_posix()}",
                    "output_json_relative_path": f"{CONVERSION_LAYER_NAME}/{json_rel.as_posix()}",
                }
            )
    finally:
        workbook.close()

    manifest_payload = {
        "project": "Vitruvius",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_relative_path": relative_source_path,
        "source_absolute_path": workbook_info["source_absolute_path"],
        "source_extension": workbook_info["source_extension"],
        "source_size_bytes": workbook_info["source_size_bytes"],
        "source_modified_utc": workbook_info["source_modified_utc"],
        "top_category": workbook_info["top_category"],
        "xlsx_workbook_path": workbook_info["xlsx_workbook_path"],
        "xlsx_workbook_relative_path": workbook_info["xlsx_workbook_relative_path"],
        "conversion_used": workbook_info["conversion_used"],
        "sheet_count": len(sheet_reports),
        "sheets": sheet_reports,
    }
    write_json(manifest_abs, manifest_payload)

    return {
        "source_relative_path": relative_source_path,
        "source_absolute_path": workbook_info["source_absolute_path"],
        "source_extension": workbook_info["source_extension"],
        "source_size_bytes": workbook_info["source_size_bytes"],
        "source_modified_utc": workbook_info["source_modified_utc"],
        "top_category": workbook_info["top_category"],
        "xlsx_workbook_path": workbook_info["xlsx_workbook_path"],
        "xlsx_workbook_relative_path": workbook_info["xlsx_workbook_relative_path"],
        "manifest_output": str(manifest_abs),
        "manifest_output_relative": f"{CONVERSION_LAYER_NAME}/{manifest_rel.as_posix()}",
        "sheet_count": len(sheet_reports),
        "sheets": sheet_reports,
        "status": "exported",
    }


def run_convert_spreadsheet_one(relative_source_path: str) -> int:
    repo_root = get_repo_root()
    registry_dir = ensure_registry_dir(repo_root)
    runs_dir = ensure_runs_dir(repo_root)
    ensure_spreadsheet_registry_exists(registry_dir)

    source_path = repo_root / RAW_LAYER_NAME / relative_source_path

    if not source_path.exists():
        print(f"ERROR: File not found: {source_path}")
        return 1

    if not is_supported_spreadsheet_file(source_path):
        print("ERROR: convert-spreadsheet-one only accepts .xlsx or .xls files.")
        return 1

    try:
        report = export_spreadsheet_workbook(
            relative_source_path=relative_source_path,
            run_command="convert-spreadsheet-one",
        )
    except Exception as exc:
        try:
            stat = source_path.stat()
            source_size_bytes = stat.st_size
            source_modified_utc = datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat()
        except FileNotFoundError:
            source_size_bytes = ""
            source_modified_utc = ""

        top_category = Path(relative_source_path).parts[0] if len(Path(relative_source_path).parts) > 0 else ""

        _write_spreadsheet_registry_entry(
            run_command="convert-spreadsheet-one",
            source_relative_path=f"{RAW_LAYER_NAME}/{Path(relative_source_path).as_posix()}",
            logical_source_relative_path=Path(relative_source_path).as_posix(),
            source_extension=source_path.suffix.lower(),
            source_size_bytes=source_size_bytes,
            source_modified_utc=source_modified_utc,
            top_category=top_category,
            sheet_name="",
            output_csv_relative_path="",
            output_json_relative_path="",
            row_count="",
            column_count="",
            status="FAILED",
            error_message=str(exc),
        )

        stamp = utc_now_stamp()
        snapshot_spreadsheet_registry(registry_dir, stamp)
        print(f"ERROR: {exc}")
        return 1

    stamp = utc_now_stamp()
    safe_stem = safe_name_for_report(source_path.stem)

    latest_json = runs_dir / "latest_convert_spreadsheet_one_report.json"
    timestamped_json = runs_dir / f"{stamp}_{safe_stem}_convert_spreadsheet_one_report.json"

    latest_txt = runs_dir / "latest_convert_spreadsheet_one_report.txt"
    timestamped_txt = runs_dir / f"{stamp}_{safe_stem}_convert_spreadsheet_one_report.txt"

    text_lines = [
        "Vitruvius Convert Spreadsheet One Report",
        "",
        f"Source: {relative_source_path}",
        f"Source extension: {report['source_extension']}",
        f"Workbook path used: {report['xlsx_workbook_path']}",
        f"Workbook relative path used: {report['xlsx_workbook_relative_path']}",
        f"Manifest output: {report['manifest_output']}",
        f"Sheet count: {report['sheet_count']}",
        "",
        "SHEETS:",
    ]

    if report["sheets"]:
        for sheet in report["sheets"]:
            text_lines.append("  ---")
            text_lines.append(f"  sheet_name: {sheet['sheet_name']}")
            text_lines.append(f"  row_count: {sheet['row_count']}")
            text_lines.append(f"  column_count: {sheet['column_count']}")
            text_lines.append(f"  csv: {sheet['output_csv_relative_path']}")
            text_lines.append(f"  json: {sheet['output_json_relative_path']}")
    else:
        text_lines.append("  None")

    report_text = "\n".join(text_lines) + "\n"

    report_payload = {
        "project": "Vitruvius",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        **report,
    }

    write_json(latest_json, report_payload)
    write_json(timestamped_json, report_payload)
    write_text(latest_txt, report_text)
    write_text(timestamped_txt, report_text)

    snapshot_spreadsheet_registry(registry_dir, stamp)

    print("Convert-spreadsheet-one complete.")
    print(f"Source file: {source_path}")
    print(f"Workbook path used: {report['xlsx_workbook_path']}")
    print(f"Manifest output: {report['manifest_output']}")
    print(f"Sheet count: {report['sheet_count']}")
    print(f"Latest spreadsheet report JSON: {latest_json}")
    print(f"Latest spreadsheet report TXT: {latest_txt}")

    return 0